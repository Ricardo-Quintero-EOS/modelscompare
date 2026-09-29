import csv
from decimal import Decimal

from openpyxl import load_workbook

from modelscompare.exporting import export_records
from modelscompare.models import ModelPricing


def _record() -> ModelPricing:
    return ModelPricing(
        name="GPT Example",
        input_cost_credits=Decimal("100"),
        output_cost_credits=Decimal("200"),
        cached_input_cost_credits=Decimal("10"),
        cache_write_cost_credits=Decimal("125"),
        usd_per_credit=Decimal("0.01"),
        provider="OpenAI",
        artificial_analysis_url="https://artificialanalysis.ai/models/gpt-example",
    )


def test_csv_and_excel_export_only_the_four_cost_types(tmp_path) -> None:
    paths = export_records([_record()], ["csv", "xlsx"], downloads_dir=tmp_path)
    csv_path, excel_path = paths

    with csv_path.open(newline="", encoding="utf-8-sig") as file:
        row = next(csv.DictReader(file))
    assert row["Input (USD / 1M token)"] == "1.00"
    assert row["Cached read (USD / 1M token)"] == "0.10"
    assert row["Cached write (USD / 1M token)"] == "1.25"
    assert row["Output (USD / 1M token)"] == "2.00"
    assert all("_" not in header for header in row)
    assert not any("blended" in header.casefold() for header in row)

    sheet = load_workbook(excel_path).active
    headers = [cell.value for cell in sheet[1]]
    input_usd_column = headers.index("Input (USD / 1M token)") + 1
    cached_read_credits_column = headers.index("Cached read (crediti AI / 1M token)") + 1
    link_column = headers.index("Scheda Artificial Analysis") + 1
    assert sheet.cell(2, input_usd_column).value == 1
    assert sheet.cell(2, input_usd_column).data_type == "n"
    assert sheet.cell(2, input_usd_column).number_format.startswith("$")
    assert sheet.cell(2, cached_read_credits_column).value == 10
    assert sheet.cell(2, cached_read_credits_column).number_format == '#,##0.00 "cr"'
    assert sheet.cell(2, link_column).hyperlink.target.endswith("/gpt-example")
    assert "Ragionamento" in headers
    assert "Finestra contesto" in headers
    assert list(sheet.tables) == ["ModelsCompare"]
    assert not any("blended" in header.casefold() for header in headers)


def test_export_respects_selected_formats_and_downloads_directory(tmp_path) -> None:
    paths = export_records([_record()], ["csv"], "nested/only_csv.xlsx", tmp_path)

    assert len(paths) == 1
    assert paths[0].name == "only_csv.csv"
    assert paths[0].parent == tmp_path


def test_export_uses_incremental_suffixes_without_overwriting_existing_files(tmp_path) -> None:
    export_records([_record()], ["csv", "xlsx"], downloads_dir=tmp_path)
    export_records([_record()], ["csv", "xlsx"], downloads_dir=tmp_path)

    paths = export_records([_record()], ["csv", "xlsx"], downloads_dir=tmp_path)

    assert [path.name for path in paths] == ["copilot_models_2.csv", "copilot_models_2.xlsx"]
    assert load_workbook(paths[1]).active["A2"].value == "GPT Example"
