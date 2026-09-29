import csv
import io
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from modelscompare.models import ModelPricing

COST_COLUMNS = (
    ("Input", "input_cost_usd", "input_cost_credits"),
    ("Cached read", "cached_input_cost_usd", "cached_input_cost_credits"),
    ("Cached write", "cache_write_cost_usd", "cache_write_cost_credits"),
    ("Output", "output_cost_usd", "output_cost_credits"),
)


def _write_unique_file(destination: Path, prefix: str, extension: str, content: bytes) -> Path:
    index = 0
    while True:
        suffix = f"_{index}" if index else ""
        path = destination / f"{prefix}{suffix}{extension}"
        try:
            file = path.open("xb")
        except FileExistsError:
            index += 1
        except PermissionError:
            if not path.exists():
                raise
            index += 1
            continue
        else:
            try:
                with file:
                    file.write(content)
                return path
            except PermissionError:
                try:
                    path.unlink()
                except OSError:
                    pass
                index += 1


def _export_columns():
    columns = [
        ("Modello", "name", "text"),
        ("Provider", "provider", "text"),
        ("Categoria", "category", "text"),
        ("Stato di rilascio", "release_status", "text"),
        ("Tier", "tier", "text"),
        ("Soglia contesto", "threshold", "text"),
        ("Ragionamento", "reasoning_label", "text"),
        ("Finestra contesto", "context_window", "text"),
    ]
    for label, usd_field, credits_field in COST_COLUMNS:
        columns.extend((
            (f"{label} (USD / 1M token)", usd_field, "currency"),
            (f"{label} (crediti AI / 1M token)", credits_field, "credits"),
        ))
    columns.append(("Scheda Artificial Analysis", "artificial_analysis_url", "link"))
    return columns


def _excel_value(record: ModelPricing, attribute: str, kind: str):
    if kind == "link":
        return "Apri scheda" if record.artificial_analysis_url else ""
    value = getattr(record, attribute)
    if isinstance(value, Decimal):
        return float(value)
    return value


def export_records(
    records: list[ModelPricing],
    formats: list[str],
    prefix: str = "copilot_models",
    downloads_dir: Path | None = None,
) -> list[Path]:
    destination = downloads_dir or Path.home() / "Downloads"
    destination.mkdir(parents=True, exist_ok=True)
    safe_prefix = Path(prefix).name or "copilot_models"
    if Path(safe_prefix).suffix.casefold() in {".csv", ".xlsx"}:
        safe_prefix = Path(safe_prefix).stem
    columns = _export_columns()
    headers = [header for header, _attribute, _kind in columns]
    output_paths: list[Path] = []

    if "csv" in formats:
        csv_buffer = io.StringIO(newline="")
        writer = csv.writer(csv_buffer)
        writer.writerow(headers)
        for record in records:
            writer.writerow([
                format(value, "f") if isinstance(value := getattr(record, attribute), Decimal) else value
                for _header, attribute, _kind in columns
            ])
        output_paths.append(_write_unique_file(
            destination, safe_prefix, ".csv", csv_buffer.getvalue().encode("utf-8-sig"),
        ))

    if "xlsx" in formats:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Copilot prices"
        sheet.sheet_properties.tabColor = "1D8792"
        sheet.append(headers)
        for record in records:
            sheet.append([_excel_value(record, attribute, kind) for _header, attribute, kind in columns])

        header_fill = PatternFill("solid", fgColor="163B4D")
        header_font = Font(name="Aptos Display", size=10, bold=True, color="FFFFFF")
        header_border = Border(bottom=Side(style="medium", color="27A6A1"))
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
            cell.border = header_border
        sheet.row_dimensions[1].height = 38

        for row_index, record in enumerate(records, start=2):
            row_fill = PatternFill("solid", fgColor="F2F7F8" if row_index % 2 == 0 else "FFFFFF")
            for column_index, (_header, _attribute, kind) in enumerate(columns, start=1):
                cell = sheet.cell(row=row_index, column=column_index)
                cell.fill = row_fill
                cell.alignment = Alignment(
                    vertical="center",
                    horizontal="right" if kind in {"currency", "credits"} else "left",
                )
                if kind == "currency" and cell.value is not None:
                    cell.number_format = '$#,##0.0000;[Red]($#,##0.0000);-'
                elif kind == "credits" and cell.value is not None:
                    cell.number_format = '#,##0.00 "cr"'
            sheet.row_dimensions[row_index].height = 21
            if record.artificial_analysis_url:
                model_cell = sheet.cell(row=row_index, column=1)
                model_cell.hyperlink = record.artificial_analysis_url
                model_cell.font = Font(color="087F8C", underline="single")
                link_cell = sheet.cell(row=row_index, column=len(columns))
                link_cell.hyperlink = record.artificial_analysis_url
                link_cell.font = Font(color="087F8C", underline="single")

        if records:
            table = Table(displayName="ModelsCompare", ref=sheet.dimensions)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            sheet.add_table(table)
        sheet.freeze_panes = "A2"
        for column_index, (header, _attribute, kind) in enumerate(columns, start=1):
            values = [
                len(str(sheet.cell(row=row_index, column=column_index).value or ""))
                for row_index in range(1, sheet.max_row + 1)
            ]
            maximum_width = 19 if kind in {"currency", "credits"} else 30
            sheet.column_dimensions[sheet.cell(row=1, column=column_index).column_letter].width = min(
                max(len(header), max(values, default=0)) + 2,
                maximum_width,
            )
        excel_buffer = io.BytesIO()
        workbook.save(excel_buffer)
        output_paths.append(_write_unique_file(
            destination, safe_prefix, ".xlsx", excel_buffer.getvalue(),
        ))

    return output_paths
